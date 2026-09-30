# -*- coding: utf-8 -*-
"""Öffnet erzeugte Office-Dateien mit unabhängigen Bibliotheken und gibt den Inhalt als JSON aus.

  python tests/referenz/pruefe_office.py datei.docx|datei.xlsx|datei.zip

- .zip/.docx/.xlsx: zipfile.testzip() prüft Prüfsummen und Aufbau des Containers
- .docx: python-docx liest Absätze, Formatvorlagen, Tabellen und eingebettete Bilder
- .xlsx: openpyxl liest Blätter, Zellwerte und Zahlenformate
Beendet sich mit Fehlercode, wenn die Datei nicht geöffnet werden kann. Auf die Standardausgabe geht
ausschließlich das JSON (siehe nur_json.py).
"""
import sys, zipfile
import nur_json


def pruefe_zip(pfad):
    with zipfile.ZipFile(pfad) as z:
        kaputt = z.testzip()
        if kaputt:
            raise SystemExit('ZIP-Prüfsumme falsch: ' + kaputt)
        return {'dateien': z.namelist()}


def pruefe_docx(pfad):
    import docx
    d = docx.Document(pfad)
    absaetze = [{'stil': p.style.name if p.style is not None else '', 'text': p.text} for p in d.paragraphs if p.text.strip()]
    tabellen = [[[c.text for c in r.cells] for r in t.rows] for t in d.tables]
    return {'absaetze': absaetze, 'tabellen': tabellen, 'bilder': len(d.inline_shapes),
            'bildgroessen_cm': [[round(s.width.cm, 2), round(s.height.cm, 2)] for s in d.inline_shapes],
            'titel': d.core_properties.title, 'sprache': d.core_properties.language}


def pruefe_xlsx(pfad):
    import openpyxl
    wb = openpyxl.load_workbook(pfad)
    blaetter = {}
    for ws in wb.worksheets:
        zellen = []
        for zeile in ws.iter_rows():
            for c in zeile:
                if c.value is not None:
                    zellen.append({'ref': c.coordinate, 'wert': c.value, 'format': c.number_format, 'fett': bool(c.font and c.font.b)})
        blaetter[ws.title] = zellen
    return {'blaetter': blaetter}


def main():
    ausgeben = nur_json.umlenken()
    pfad = sys.argv[1]
    erg = {'zip': pruefe_zip(pfad)}
    if pfad.endswith('.docx'):
        erg['docx'] = pruefe_docx(pfad)
    elif pfad.endswith('.xlsx'):
        erg['xlsx'] = pruefe_xlsx(pfad)
    ausgeben(erg)


if __name__ == '__main__':
    main()
